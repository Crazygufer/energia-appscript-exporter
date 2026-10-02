from datetime import datetime, timedelta, time
from pathlib import Path
from typing import Optional

import pandas as pd
from openpyxl import Workbook

from .config import (
    COLUMNAS_DGA,
    DGA_FLUJO_INSTANTANEO_VAR,
    DGA_HORA_DEFAULT,
    DGA_HORA_LIMITE_BUSQUEDA,
    DGA_NIVEL_FREATICO_VAR,
    DGA_TOTALIZADOR_VAR,
    TZ_DGA,
)

from .database import (
    obtener_datos_por_rango_exclusivo,
)


# ============================================================
# NORMALIZACIÓN
# ============================================================

def _normalizar_df(
    df: pd.DataFrame,
) -> pd.DataFrame:
    """
    Normaliza timestamp/valor y convierte el timestamp Unix
    a la zona horaria fija utilizada por el reporte DGA.
    """

    if df.empty:
        return pd.DataFrame(
            columns=[
                "timestamp",
                "valor",
                "fecha_hora_dga",
            ]
        )

    df = df.copy()

    df["timestamp"] = pd.to_numeric(
        df["timestamp"],
        errors="coerce",
    )

    df["valor"] = pd.to_numeric(
        df["valor"],
        errors="coerce",
    )

    df = df.dropna(
        subset=[
            "timestamp",
            "valor",
        ]
    )

    if df.empty:
        return pd.DataFrame(
            columns=[
                "timestamp",
                "valor",
                "fecha_hora_dga",
            ]
        )

    df["fecha_hora_dga"] = (
        pd.to_datetime(
            df["timestamp"],
            unit="s",
            utc=True,
        )
        .dt.tz_convert(TZ_DGA)
    )

    return (
        df
        .sort_values(
            "fecha_hora_dga"
        )
        .reset_index(
            drop=True
        )
    )


def _df_vacio() -> pd.DataFrame:
    """
    Devuelve un DataFrame compatible con los datos
    normalizados del reporte DGA.
    """

    return pd.DataFrame(
        columns=[
            "timestamp",
            "valor",
            "fecha_hora_dga",
        ]
    )


# ============================================================
# BÚSQUEDA EN VENTANA DGA
# ============================================================

def _buscar_primera_muestra_en_ventana(
    df: pd.DataFrame,
    fecha_base: datetime,
    hora_inicio: int,
    hora_fin: int,
) -> Optional[pd.Series]:
    """
    Busca la PRIMERA lectura disponible del día entre
    hora_inicio y hora_fin.

    Ejemplo:

        hora oficial:
            07:00

        ventana:
            07:00 -> 11:00

        registros:
            06:59 -> no se usa
            08:49 -> se usa
            09:04 -> no se usa

    El Excel oficial seguirá informando 07:00.
    """

    if df.empty:
        return None

    inicio = fecha_base.replace(
        hour=hora_inicio,
        minute=0,
        second=0,
        microsecond=0,
        tzinfo=TZ_DGA,
    )

    fin = fecha_base.replace(
        hour=hora_fin,
        minute=0,
        second=0,
        microsecond=0,
        tzinfo=TZ_DGA,
    )

    if hora_fin < hora_inicio:
        raise ValueError(
            "La hora límite DGA no puede ser "
            "menor que la hora inicial."
        )

    candidatos = df[
        (
            df["fecha_hora_dga"]
            >= inicio
        )
        &
        (
            df["fecha_hora_dga"]
            <= fin
        )
    ]

    if candidatos.empty:
        return None

    return candidatos.iloc[0]


# ============================================================
# UTILIDADES DE VALOR
# ============================================================

def _obtener_valor_muestra(
    muestra: Optional[pd.Series],
    decimales: int = 3,
) -> Optional[float]:
    """
    Obtiene el valor numérico de una muestra.
    """

    if muestra is None:
        return None

    return round(
        float(
            muestra["valor"]
        ),
        decimales,
    )


def _obtener_timestamp_muestra(
    muestra: Optional[pd.Series],
) -> str:
    """
    Devuelve el timestamp real utilizado en formato legible.
    """

    if muestra is None:
        return ""

    return muestra[
        "fecha_hora_dga"
    ].strftime(
        "%Y-%m-%d %H:%M:%S"
    )


def _obtener_hora_muestra(
    muestra: Optional[pd.Series],
) -> str:
    """
    Devuelve solamente la hora real de la muestra.
    """

    if muestra is None:
        return ""

    return muestra[
        "fecha_hora_dga"
    ].strftime(
        "%H:%M:%S"
    )


# ============================================================
# CÁLCULO DEL CAUDAL MEDIO
# ============================================================

def _calcular_caudal_medio_periodo_lps(
    muestra_actual: pd.Series,
    muestra_anterior: Optional[pd.Series],
) -> Optional[float]:
    """
    Calcula el caudal promedio utilizando el incremento
    real del totalizador y el tiempo realmente transcurrido.

        Q [l/s] =
            (Totalizador_actual - Totalizador_anterior) * 1000
            -------------------------------------------------
                       segundos transcurridos

    Ejemplo:

        Día anterior:
            07:03

        Día actual:
            08:49

    No se fuerza una diferencia artificial de 24 horas.
    Se utiliza el tiempo real entre ambas lecturas.
    """

    if muestra_anterior is None:
        return None

    valor_actual = float(
        muestra_actual["valor"]
    )

    valor_anterior = float(
        muestra_anterior["valor"]
    )

    delta_m3 = (
        valor_actual
        - valor_anterior
    )

    # Reset, rollover o anomalía del totalizador.
    if delta_m3 < 0:
        return None

    delta_s = (
        muestra_actual[
            "fecha_hora_dga"
        ]
        -
        muestra_anterior[
            "fecha_hora_dga"
        ]
    ).total_seconds()

    if delta_s <= 0:
        return None

    return round(
        (
            delta_m3
            * 1000.0
        )
        / delta_s,
        3,
    )


# ============================================================
# GENERACIÓN DEL REPORTE
# ============================================================

def generar_reporte_dga(
    fecha_inicio: str,
    fecha_fin: str,
    hora_medicion: int = DGA_HORA_DEFAULT,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """
    Genera el reporte diario del pozo.

    Variables:

        Totalizado_Pozo
            Lectura acumulada del medidor.

        Nivel_Pozo
            Lectura real del nivel freático.

        Flujo_Instantaneo
            Lectura instantánea informada por el instrumento.

        Caudal (l/s)
            Se calcula independientemente a partir de la
            diferencia del totalizador entre dos muestras.

    Regla temporal:

        Hora oficial:
            07:00

        Ventana de búsqueda:
            07:00 -> 11:00

        Selección:
            Primera muestra válida encontrada.

    La hora oficial informada en el archivo DGA continúa
    siendo 07:00 aunque la lectura real haya ocurrido
    posteriormente.
    """

    hora_medicion = int(
        hora_medicion
    )

    hora_limite = int(
        DGA_HORA_LIMITE_BUSQUEDA
    )

    if not 0 <= hora_medicion <= 23:
        raise ValueError(
            "La hora DGA debe estar entre 0 y 23."
        )

    if not 0 <= hora_limite <= 23:
        raise ValueError(
            "La hora límite DGA debe estar "
            "entre 0 y 23."
        )

    if hora_limite < hora_medicion:
        raise ValueError(
            "La hora límite DGA debe ser igual "
            "o posterior a la hora oficial."
        )


    # --------------------------------------------------------
    # Fechas
    # --------------------------------------------------------

    fecha_ini = datetime.strptime(
        fecha_inicio,
        "%Y-%m-%d",
    )

    fecha_ter = datetime.strptime(
        fecha_fin,
        "%Y-%m-%d",
    )

    if fecha_ini > fecha_ter:
        raise ValueError(
            "La fecha inicial no puede ser posterior "
            "a la fecha final."
        )


    # ========================================================
    # RANGO DE CONSULTA
    # ========================================================

    # Se consulta un día anterior para poder calcular
    # el caudal del primer día solicitado.
    consulta_inicio = (
        fecha_ini.replace(
            hour=hora_medicion,
            minute=0,
            second=0,
            microsecond=0,
            tzinfo=TZ_DGA,
        )
        - timedelta(days=1)
    )

    # Se agrega un minuto porque la consulta SQL utiliza:
    #
    #     [inicio, fin)
    #
    # y necesitamos incluir exactamente la lectura de las 11:00.
    consulta_fin = (
        fecha_ter.replace(
            hour=hora_limite,
            minute=0,
            second=0,
            microsecond=0,
            tzinfo=TZ_DGA,
        )
        + timedelta(minutes=1)
    )


    # ========================================================
    # TOTALIZADOR
    # ========================================================

    df_totalizador = _normalizar_df(
        obtener_datos_por_rango_exclusivo(
            DGA_TOTALIZADOR_VAR,
            consulta_inicio,
            consulta_fin,
        )
    )


    # ========================================================
    # NIVEL FREÁTICO
    # ========================================================

    if DGA_NIVEL_FREATICO_VAR:

        df_nivel = _normalizar_df(
            obtener_datos_por_rango_exclusivo(
                DGA_NIVEL_FREATICO_VAR,
                consulta_inicio,
                consulta_fin,
            )
        )

    else:

        df_nivel = _df_vacio()


    # ========================================================
    # FLUJO INSTANTÁNEO
    # ========================================================

    if DGA_FLUJO_INSTANTANEO_VAR:

        df_flujo = _normalizar_df(
            obtener_datos_por_rango_exclusivo(
                DGA_FLUJO_INSTANTANEO_VAR,
                consulta_inicio,
                consulta_fin,
            )
        )

    else:

        df_flujo = _df_vacio()


    # ========================================================
    # GENERACIÓN DE FILAS
    # ========================================================

    filas_dga = []

    filas_traza = []

    fecha_actual = fecha_ini


    while fecha_actual <= fecha_ter:

        # ----------------------------------------------------
        # TOTALIZADOR ACTUAL
        # ----------------------------------------------------

        muestra_actual = (
            _buscar_primera_muestra_en_ventana(
                df_totalizador,
                fecha_actual,
                hora_medicion,
                hora_limite,
            )
        )


        # ----------------------------------------------------
        # TOTALIZADOR DÍA ANTERIOR
        # ----------------------------------------------------

        fecha_anterior = (
            fecha_actual
            - timedelta(days=1)
        )

        muestra_anterior = (
            _buscar_primera_muestra_en_ventana(
                df_totalizador,
                fecha_anterior,
                hora_medicion,
                hora_limite,
            )
        )


        # ----------------------------------------------------
        # NIVEL FREÁTICO
        # ----------------------------------------------------

        muestra_nivel = (
            _buscar_primera_muestra_en_ventana(
                df_nivel,
                fecha_actual,
                hora_medicion,
                hora_limite,
            )
        )

        nivel = _obtener_valor_muestra(
            muestra_nivel,
            decimales=3,
        )


        # ----------------------------------------------------
        # FLUJO INSTANTÁNEO
        # ----------------------------------------------------

        muestra_flujo = (
            _buscar_primera_muestra_en_ventana(
                df_flujo,
                fecha_actual,
                hora_medicion,
                hora_limite,
            )
        )

        flujo_instantaneo = (
            _obtener_valor_muestra(
                muestra_flujo,
                decimales=6,
            )
        )


        # ====================================================
        # DÍA SIN TOTALIZADOR
        # ====================================================

        if muestra_actual is None:

            filas_dga.append({
                "Fecha Medición":
                    fecha_actual.date(),

                "Hora Medición":
                    time(
                        hour=hora_medicion
                    ),

                "Caudal (l/s)":
                    "",

                "Totalizador (m3)":
                    "",

                "Nivel Freático (m)":
                    (
                        ""
                        if nivel is None
                        else nivel
                    ),

                "Flujo Instantáneo":
                    (
                        ""
                        if flujo_instantaneo is None
                        else flujo_instantaneo
                    ),
            })


            estado = (
                "SIN TOTALIZADOR ENTRE "
                f"{hora_medicion:02d}:00 Y "
                f"{hora_limite:02d}:00"
            )

            if (
                DGA_NIVEL_FREATICO_VAR
                and nivel is None
            ):
                estado += " / SIN NIVEL"

            if (
                DGA_FLUJO_INSTANTANEO_VAR
                and flujo_instantaneo is None
            ):
                estado += " / SIN FLUJO"


            filas_traza.append({
                "fecha":
                    fecha_actual.strftime(
                        "%Y-%m-%d"
                    ),

                "hora_reportada":
                    f"{hora_medicion:02d}:00",

                "ventana_busqueda":
                    (
                        f"{hora_medicion:02d}:00-"
                        f"{hora_limite:02d}:00"
                    ),

                "timestamp_real":
                    "",

                "hora_real":
                    "",

                "retraso_min":
                    "",

                "totalizador_m3":
                    "",

                "totalizador_anterior_m3":
                    "",

                "timestamp_anterior":
                    "",

                "caudal_lps":
                    "",

                "nivel_freatico_m":
                    (
                        ""
                        if nivel is None
                        else nivel
                    ),

                "timestamp_nivel":
                    _obtener_timestamp_muestra(
                        muestra_nivel
                    ),

                "hora_real_nivel":
                    _obtener_hora_muestra(
                        muestra_nivel
                    ),

                "flujo_instantaneo":
                    (
                        ""
                        if flujo_instantaneo is None
                        else flujo_instantaneo
                    ),

                "timestamp_flujo":
                    _obtener_timestamp_muestra(
                        muestra_flujo
                    ),

                "hora_real_flujo":
                    _obtener_hora_muestra(
                        muestra_flujo
                    ),

                "estado":
                    estado,
            })


            fecha_actual += timedelta(
                days=1
            )

            continue


        # ====================================================
        # TOTALIZADOR
        # ====================================================

        totalizador = round(
            float(
                muestra_actual[
                    "valor"
                ]
            ),
            3,
        )


        # ====================================================
        # CAUDAL PROMEDIO CALCULADO
        # ====================================================

        caudal = (
            _calcular_caudal_medio_periodo_lps(
                muestra_actual,
                muestra_anterior,
            )
        )


        # ====================================================
        # ESTADO
        # ====================================================

        if muestra_anterior is None:

            estado = (
                "OK TOTALIZADOR / "
                "SIN DIA ANTERIOR PARA CAUDAL"
            )

        elif caudal is None:

            estado = (
                "REVISAR CAUDAL / TOTALIZADOR"
            )

        else:

            estado = "OK"


        if (
            DGA_NIVEL_FREATICO_VAR
            and nivel is None
        ):

            estado += " / SIN NIVEL"


        if (
            DGA_FLUJO_INSTANTANEO_VAR
            and flujo_instantaneo is None
        ):

            estado += " / SIN FLUJO"


        # ====================================================
        # ARCHIVO OFICIAL
        # ====================================================

        filas_dga.append({
            "Fecha Medición":
                fecha_actual.date(),

            # Hora oficial reportada.
            "Hora Medición":
                time(
                    hour=hora_medicion
                ),

            "Caudal (l/s)":
                (
                    ""
                    if caudal is None
                    else caudal
                ),

            "Totalizador (m3)":
                totalizador,

            "Nivel Freático (m)":
                (
                    ""
                    if nivel is None
                    else nivel
                ),

            "Flujo Instantáneo":
                (
                    ""
                    if flujo_instantaneo is None
                    else flujo_instantaneo
                ),
        })


        # ====================================================
        # TRAZABILIDAD
        # ====================================================

        hora_oficial_dt = (
            fecha_actual.replace(
                hour=hora_medicion,
                minute=0,
                second=0,
                microsecond=0,
                tzinfo=TZ_DGA,
            )
        )


        retraso_min = round(
            (
                muestra_actual[
                    "fecha_hora_dga"
                ]
                - hora_oficial_dt
            ).total_seconds()
            / 60.0,
            2,
        )


        filas_traza.append({
            "fecha":
                fecha_actual.strftime(
                    "%Y-%m-%d"
                ),

            "hora_reportada":
                f"{hora_medicion:02d}:00",

            "ventana_busqueda":
                (
                    f"{hora_medicion:02d}:00-"
                    f"{hora_limite:02d}:00"
                ),

            "timestamp_real":
                _obtener_timestamp_muestra(
                    muestra_actual
                ),

            "hora_real":
                _obtener_hora_muestra(
                    muestra_actual
                ),

            "retraso_min":
                retraso_min,

            "totalizador_m3":
                totalizador,

            "totalizador_anterior_m3":
                (
                    ""
                    if muestra_anterior is None
                    else round(
                        float(
                            muestra_anterior[
                                "valor"
                            ]
                        ),
                        3,
                    )
                ),

            "timestamp_anterior":
                _obtener_timestamp_muestra(
                    muestra_anterior
                ),

            "caudal_lps":
                (
                    ""
                    if caudal is None
                    else caudal
                ),

            "nivel_freatico_m":
                (
                    ""
                    if nivel is None
                    else nivel
                ),

            "timestamp_nivel":
                _obtener_timestamp_muestra(
                    muestra_nivel
                ),

            "hora_real_nivel":
                _obtener_hora_muestra(
                    muestra_nivel
                ),

            "flujo_instantaneo":
                (
                    ""
                    if flujo_instantaneo is None
                    else flujo_instantaneo
                ),

            "timestamp_flujo":
                _obtener_timestamp_muestra(
                    muestra_flujo
                ),

            "hora_real_flujo":
                _obtener_hora_muestra(
                    muestra_flujo
                ),

            "estado":
                estado,
        })


        fecha_actual += timedelta(
            days=1
        )


    # ========================================================
    # DATAFRAMES FINALES
    # ========================================================

    df_dga = pd.DataFrame(
        filas_dga,
        columns=COLUMNAS_DGA,
    )

    df_traza = pd.DataFrame(
        filas_traza
    )

    return (
        df_dga,
        df_traza,
    )


# ============================================================
# GENERACIÓN DEL EXCEL
# ============================================================

def _guardar_excel_plano(
    df: pd.DataFrame,
    path: Path,
    sheet_name: str = "Datos",
    formato_dga: bool = False,
) -> None:
    """
    Escribe un archivo XLSX.
    """

    wb = Workbook()

    ws = wb.active

    ws.title = sheet_name


    # --------------------------------------------------------
    # Encabezados
    # --------------------------------------------------------

    ws.append(
        list(
            df.columns
        )
    )


    # --------------------------------------------------------
    # Datos
    # --------------------------------------------------------

    for fila in df.itertuples(
        index=False,
        name=None,
    ):

        ws.append(
            list(
                fila
            )
        )


    # --------------------------------------------------------
    # Formato DGA
    # --------------------------------------------------------

    if formato_dga:

        for row in range(
            2,
            ws.max_row + 1,
        ):

            ws[
                f"A{row}"
            ].number_format = (
                "dd/mm/yyyy"
            )

            ws[
                f"B{row}"
            ].number_format = (
                "hh:mm"
            )

            ws[
                f"C{row}"
            ].number_format = (
                "0.000"
            )

            ws[
                f"D{row}"
            ].number_format = (
                "0.000"
            )

            ws[
                f"E{row}"
            ].number_format = (
                "0.000"
            )

            if ws.max_column >= 6:

                ws[
                    f"F{row}"
                ].number_format = (
                    "0.000000"
                )


        ws.column_dimensions[
            "A"
        ].width = 18

        ws.column_dimensions[
            "B"
        ].width = 16

        ws.column_dimensions[
            "C"
        ].width = 16

        ws.column_dimensions[
            "D"
        ].width = 18

        ws.column_dimensions[
            "E"
        ].width = 20

        ws.column_dimensions[
            "F"
        ].width = 20


    wb.save(
        path
    )


# ============================================================
# GUARDAR REPORTE DGA
# ============================================================

def guardar_reporte_dga_excel(
    path: Path,
    fecha_inicio: str,
    fecha_fin: str,
    hora_medicion: int = DGA_HORA_DEFAULT,
) -> dict:
    """
    Guarda:

        - archivo oficial DGA;
        - archivo de trazabilidad técnica.

    Además entrega estadísticas sobre la cantidad de
    días que tienen datos disponibles.
    """

    path = Path(
        path
    )


    path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )


    df_dga, df_traza = (
        generar_reporte_dga(
            fecha_inicio=fecha_inicio,
            fecha_fin=fecha_fin,
            hora_medicion=hora_medicion,
        )
    )


    # --------------------------------------------------------
    # Archivo oficial
    # --------------------------------------------------------

    _guardar_excel_plano(
        df_dga,
        path,
        sheet_name="Datos",
        formato_dga=True,
    )


    # --------------------------------------------------------
    # Archivo de trazabilidad
    # --------------------------------------------------------

    trazabilidad_path = (
        path.with_name(
            f"{path.stem}_trazabilidad.xlsx"
        )
    )


    _guardar_excel_plano(
        df_traza,
        trazabilidad_path,
        sheet_name="Trazabilidad",
        formato_dga=False,
    )


    # ========================================================
    # ESTADÍSTICAS
    # ========================================================

    total_dias = len(
        df_dga
    )


    con_totalizador = int(
        pd.to_numeric(
            df_dga[
                "Totalizador (m3)"
            ],
            errors="coerce",
        )
        .notna()
        .sum()
    )


    con_caudal = int(
        pd.to_numeric(
            df_dga[
                "Caudal (l/s)"
            ],
            errors="coerce",
        )
        .notna()
        .sum()
    )


    con_nivel = int(
        pd.to_numeric(
            df_dga[
                "Nivel Freático (m)"
            ],
            errors="coerce",
        )
        .notna()
        .sum()
    )


    con_flujo = int(
        pd.to_numeric(
            df_dga[
                "Flujo Instantáneo"
            ],
            errors="coerce",
        )
        .notna()
        .sum()
    )


    return {
        "path":
            path,

        "trazabilidad_path":
            trazabilidad_path,

        "rows":
            total_dias,

        "con_totalizador":
            con_totalizador,

        "con_caudal":
            con_caudal,

        "con_nivel":
            con_nivel,

        "con_flujo":
            con_flujo,

        "hora_medicion":
            int(
                hora_medicion
            ),

        "hora_limite_busqueda":
            DGA_HORA_LIMITE_BUSQUEDA,
    }