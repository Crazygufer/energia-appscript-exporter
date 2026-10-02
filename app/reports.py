from datetime import datetime
from pathlib import Path
import re
import shutil
from typing import Optional

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import (
    COLUMNAS_REPORTE_AGUA,
    COLUMNAS_TENDENCIA_AGUA,
    COLUMNAS_RESUMEN_DIARIO,
    DEFAULT_OUTPUT_XLSX,
    DEFAULT_OUTPUT_RESUMEN_DIARIO_XLSX,
    MESES_ES,
    MESES_ES_COMPLETOS,
    MESES_HISTORICOS,
    OPERATIONAL_DAY_START_HOUR,
    REGLA_DIA_OPERATIVO,
    SHEET_INDICE_MESES,
    SHEET_RESUMEN_AGUA,
    SHEET_TENDENCIA_AGUA,
    SHEET_RESUMEN_DIARIO,
    TZ_CHILE,
    VARIABLE_AGUA,
)
from .database import obtener_datos_por_rango_exclusivo
from .time_ranges import (
    add_months,
    build_agua_ranges,
    format_day_label,
    get_operational_month_window,
    get_operational_date,
    iter_day_starts,
    iter_month_starts,
    iter_week_starts,
    start_of_month_operational,
    start_of_year_operational,
)


def calcular_consumo_variable_rango_detalle(
    nombre_variable: str,
    inicio_dt: datetime,
    fin_dt: datetime,
) -> dict:
    """Calcula consumo y trazabilidad dentro del intervalo [inicio, fin)."""
    resultado_vacio = {
        "agua_m3": 0.0,
        "valor_inicial": "",
        "valor_final": "",
        "registros": 0,
    }

    if inicio_dt >= fin_dt:
        return resultado_vacio.copy()

    df = obtener_datos_por_rango_exclusivo(
        nombre_variable,
        inicio_dt,
        fin_dt,
    )

    if df.empty:
        return resultado_vacio.copy()

    df["timestamp"] = pd.to_numeric(df["timestamp"], errors="coerce")
    df["valor"] = pd.to_numeric(df["valor"], errors="coerce")
    df = df.dropna(subset=["timestamp", "valor"]).sort_values("timestamp")

    if len(df) < 2:
        resultado = resultado_vacio.copy()
        resultado["registros"] = len(df)
        return resultado

    valor_inicial = float(df.iloc[0]["valor"])
    valor_final = float(df.iloc[-1]["valor"])
    consumo = max(valor_final - valor_inicial, 0.0)

    return {
        "agua_m3": round(consumo, 3),
        "valor_inicial": round(valor_inicial, 3),
        "valor_final": round(valor_final, 3),
        "registros": int(len(df)),
    }


def calcular_resumen_agua(
    now: Optional[datetime] = None,
) -> pd.DataFrame:
    ranges = build_agua_ranges(now)
    rows = []
    fecha_actualizacion = datetime.now(TZ_CHILE).strftime("%Y-%m-%d %H:%M:%S")

    for bloque, info in ranges.items():
        detalle = calcular_consumo_variable_rango_detalle(
            VARIABLE_AGUA,
            info["from"],
            info["to"],
        )

        rows.append({
            "fecha_actualizacion": fecha_actualizacion,
            "origen": "PythonLocal",
            "regla_dia_operativo": REGLA_DIA_OPERATIVO,
            "bloque": bloque,
            "periodo": info["periodo"],
            "agua_m3": detalle["agua_m3"],
            "variable": VARIABLE_AGUA,
            "valor_inicial": detalle["valor_inicial"],
            "valor_final": detalle["valor_final"],
            "registros": detalle["registros"],
            "desde": info["from"].strftime("%Y-%m-%d %H:%M:%S"),
            "hasta": info["to"].strftime("%Y-%m-%d %H:%M:%S"),
        })

    return pd.DataFrame(rows, columns=COLUMNAS_REPORTE_AGUA)


def calcular_consumo_tendencia_rango(
    nombre_variable: str,
    inicio_dt: datetime,
    fin_dt: datetime,
) -> dict:
    detalle = calcular_consumo_variable_rango_detalle(
        nombre_variable,
        inicio_dt,
        fin_dt,
    )

    agua = float(detalle.get("agua_m3") or 0)

    return {
        "agua_m3": round(agua, 3),
        "production_tons": 0.0,
        "kpi_m3_ton": 0.0,
        "valor_inicial": detalle.get("valor_inicial", ""),
        "valor_final": detalle.get("valor_final", ""),
        "registros": int(detalle.get("registros") or 0),
    }


def crear_fila_tendencia(
    *,
    fecha_actualizacion: str,
    vista: str,
    bloque: str,
    periodo: str,
    orden: int,
    inicio_dt: datetime,
    fin_dt: datetime,
) -> dict:
    detalle = calcular_consumo_tendencia_rango(
        VARIABLE_AGUA,
        inicio_dt,
        fin_dt,
    )

    return {
        "fecha_actualizacion": fecha_actualizacion,
        "origen": "PythonLocal",
        "regla_dia_operativo": REGLA_DIA_OPERATIVO,
        "vista": vista,
        "bloque": bloque,
        "periodo": periodo,
        "orden": int(orden),
        "agua_m3": detalle["agua_m3"],
        "production_tons": detalle["production_tons"],
        "kpi_m3_ton": detalle["kpi_m3_ton"],
        "variable": VARIABLE_AGUA,
        "valor_inicial": detalle["valor_inicial"],
        "valor_final": detalle["valor_final"],
        "registros": detalle["registros"],
        "desde": inicio_dt.strftime("%Y-%m-%d %H:%M:%S"),
        "hasta": fin_dt.strftime("%Y-%m-%d %H:%M:%S"),
    }


def calcular_tendencia_agua(
    now: Optional[datetime] = None,
    months_back: int = 0,
) -> pd.DataFrame:
    """
    Genera la tendencia de un mes operacional.

    months_back = 0 genera el mes actual hasta el último corte cerrado.
    months_back = 1 genera el mes anterior completo.
    months_back = 2 genera el mes anterior anterior completo, etc.

    Las columnas y el orden de las vistas se mantienen iguales a
    Agua_Tendencia para que cada hoja histórica sea intercambiable.
    """
    if now is None:
        now = datetime.now(TZ_CHILE)

    if months_back < 0:
        raise ValueError("months_back no puede ser negativo.")

    target_month_start, target_month_end = get_operational_month_window(
        now=now,
        months_back=months_back,
    )
    year_start = start_of_year_operational(target_month_start)

    fecha_actualizacion = datetime.now(TZ_CHILE).strftime("%Y-%m-%d %H:%M:%S")
    rows = []

    # 1) Mes acumulado: consumo mensual desde enero hasta el mes objetivo.
    for orden, (inicio, fin) in enumerate(
        iter_month_starts(year_start, target_month_end),
        start=1,
    ):
        if inicio < fin:
            rows.append(crear_fila_tendencia(
                fecha_actualizacion=fecha_actualizacion,
                vista="mes_acumulado",
                bloque=inicio.strftime("%Y-%m"),
                periodo=MESES_ES.get(inicio.month, inicio.strftime("%m")),
                orden=orden,
                inicio_dt=inicio,
                fin_dt=fin,
            ))

    # 2) Semanas del mes objetivo.
    orden = 1
    for inicio, fin, week_start in iter_week_starts(
        target_month_start,
        target_month_end,
    ):
        if inicio >= fin:
            continue

        iso = week_start.isocalendar()
        rows.append(crear_fila_tendencia(
            fecha_actualizacion=fecha_actualizacion,
            vista="semanas_mes",
            bloque=f"{iso.year}-W{iso.week:02d}",
            periodo=f"Sem {iso.week:02d}",
            orden=orden,
            inicio_dt=inicio,
            fin_dt=fin,
        ))
        orden += 1

    # 3) Semanas ISO del año hasta el cierre del mes objetivo.
    orden = 1
    for inicio, fin, week_start in iter_week_starts(
        year_start,
        target_month_end,
    ):
        if inicio >= fin:
            continue

        iso = week_start.isocalendar()
        rows.append(crear_fila_tendencia(
            fecha_actualizacion=fecha_actualizacion,
            vista="semana_anio",
            bloque=f"{iso.year}-W{iso.week:02d}",
            periodo=f"{iso.year}-W{iso.week:02d}",
            orden=orden,
            inicio_dt=inicio,
            fin_dt=fin,
        ))
        orden += 1

    # 4) Días operacionales del mes objetivo.
    orden = 1
    for inicio, fin in iter_day_starts(
        target_month_start,
        target_month_end,
    ):
        if inicio >= fin:
            continue

        rows.append(crear_fila_tendencia(
            fecha_actualizacion=fecha_actualizacion,
            vista="dia_operacional",
            bloque=inicio.strftime("%Y-%m-%d"),
            periodo=format_day_label(inicio),
            orden=orden,
            inicio_dt=inicio,
            fin_dt=fin,
        ))
        orden += 1

    # 5) Año acumulado: acumulado progresivo hasta cada cierre mensual.
    orden = 1
    for inicio_mes, fin_mes in iter_month_starts(
        year_start,
        target_month_end,
    ):
        if year_start >= fin_mes:
            continue

        rows.append(crear_fila_tendencia(
            fecha_actualizacion=fecha_actualizacion,
            vista="anio_acumulado",
            bloque=inicio_mes.strftime("%Y-%m"),
            periodo=MESES_ES.get(inicio_mes.month, inicio_mes.strftime("%m")),
            orden=orden,
            inicio_dt=year_start,
            fin_dt=fin_mes,
        ))
        orden += 1

    return pd.DataFrame(rows, columns=COLUMNAS_TENDENCIA_AGUA)


def _ajustar_hoja_excel(worksheet, dataframe: pd.DataFrame) -> None:
    """
    Aplica formato visual sin modificar datos, columnas ni orden de filas.
    """
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="007CBA",
    )
    header_font = Font(
        bold=True,
        color="FFFFFF",
    )

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    for column_index, column_name in enumerate(dataframe.columns, start=1):
        values = [str(column_name)]
        values.extend(
            "" if value is None else str(value)
            for value in dataframe.iloc[:, column_index - 1].tolist()
        )

        max_length = max((len(value) for value in values), default=10)
        width = min(max(max_length + 2, 12), 32)
        worksheet.column_dimensions[get_column_letter(column_index)].width = width

    worksheet.row_dimensions[1].height = 22




# ============================================================
# RESUMEN DIARIO EXCEL POR RANGO - VARIABLE SELECCIONADA
# ============================================================

def _parse_fecha_operacional_inicio(fecha_txt: str) -> datetime:
    """Convierte AAAA-MM-DD al corte operacional 07:00 local."""
    return datetime.strptime(fecha_txt.strip(), "%Y-%m-%d").replace(
        hour=OPERATIONAL_DAY_START_HOUR,
        minute=0,
        second=0,
        microsecond=0,
        tzinfo=TZ_CHILE,
    )


def _iter_rangos_diarios_operacionales(
    fecha_inicio: str,
    fecha_fin: str,
):
    """
    Itera días operacionales [07:00, 07:00) usando fechas inclusivas.

    Ejemplo:
        fecha_inicio = 2026-09-01
        fecha_fin    = 2026-09-30

    Genera:
        2026-09-01 07:00 -> 2026-09-02 07:00
        ...
        2026-09-30 07:00 -> 2026-10-01 07:00
    """
    inicio = _parse_fecha_operacional_inicio(fecha_inicio)
    fin_ultimo_dia = _parse_fecha_operacional_inicio(fecha_fin)
    fin_exclusivo = fin_ultimo_dia + pd.Timedelta(days=1)

    if inicio >= fin_exclusivo:
        raise ValueError("La fecha inicial debe ser menor o igual a la fecha final.")

    actual = inicio
    while actual < fin_exclusivo:
        siguiente = actual + pd.Timedelta(days=1)
        yield actual, siguiente
        actual = siguiente


def _calcular_consumo_diario_generico(
    nombre_variable: str,
    inicio_dt: datetime,
    fin_dt: datetime,
) -> dict:
    """Calcula consumo por resta de acumulados dentro de [inicio, fin)."""
    df = obtener_datos_por_rango_exclusivo(
        nombre_variable,
        inicio_dt,
        fin_dt,
    )

    if df.empty:
        return {
            "consumo": "",
            "valor_inicial": "",
            "valor_final": "",
        }

    df = df.copy()
    df["timestamp"] = pd.to_numeric(df["timestamp"], errors="coerce")
    df["valor"] = pd.to_numeric(df["valor"], errors="coerce")
    df = df.dropna(subset=["timestamp", "valor"]).sort_values("timestamp")

    if df.empty:
        return {
            "consumo": "",
            "valor_inicial": "",
            "valor_final": "",
        }

    valor_inicial = float(df.iloc[0]["valor"])
    valor_final = float(df.iloc[-1]["valor"])

    if len(df) < 2:
        return {
            "consumo": 0,
            "valor_inicial": round(valor_inicial, 3),
            "valor_final": round(valor_final, 3),
        }

    consumo = valor_final - valor_inicial

    if consumo < 0:
        return {
            "consumo": "",
            "valor_inicial": round(valor_inicial, 3),
            "valor_final": round(valor_final, 3),
        }

    return {
        "consumo": round(consumo, 3),
        "valor_inicial": round(valor_inicial, 3),
        "valor_final": round(valor_final, 3),
    }


def calcular_resumen_diario_variable(
    nombre_variable: str,
    fecha_inicio: str,
    fecha_fin: str,
) -> pd.DataFrame:
    """
    Genera resumen diario 07:00-07:00 para una variable.

    Columnas finales:
        fecha_operacional, variable, consumo, valor_inicial, valor_final
    """
    rows = []

    for inicio_dt, fin_dt in _iter_rangos_diarios_operacionales(
        fecha_inicio,
        fecha_fin,
    ):
        detalle = _calcular_consumo_diario_generico(
            nombre_variable,
            inicio_dt,
            fin_dt,
        )

        rows.append({
            "fecha_operacional": inicio_dt.strftime("%Y-%m-%d"),
            "variable": nombre_variable,
            "consumo": detalle["consumo"],
            "valor_inicial": detalle["valor_inicial"],
            "valor_final": detalle["valor_final"],
        })

    return pd.DataFrame(rows, columns=COLUMNAS_RESUMEN_DIARIO)


def guardar_resumen_diario_variable_excel(
    path: Path = DEFAULT_OUTPUT_RESUMEN_DIARIO_XLSX,
    *,
    nombre_variable: str,
    fecha_inicio: str,
    fecha_fin: str,
) -> dict:
    """Guarda el resumen diario en un Excel simple."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    df = calcular_resumen_diario_variable(
        nombre_variable=nombre_variable,
        fecha_inicio=fecha_inicio,
        fecha_fin=fecha_fin,
    )

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(
            writer,
            sheet_name=SHEET_RESUMEN_DIARIO,
            index=False,
        )
        _ajustar_hoja_excel(writer.book[SHEET_RESUMEN_DIARIO], df)

    total_consumo = pd.to_numeric(df["consumo"], errors="coerce").sum()

    return {
        "path": path,
        "sheet": SHEET_RESUMEN_DIARIO,
        "rows": len(df),
        "total_consumo": round(float(total_consumo), 3),
    }


def _nombre_hoja_mes(inicio_mes: datetime) -> str:
    """
    Construye un nombre estable y legible, por ejemplo Agosto_2026.

    El año evita colisiones cuando el libro acumula varios ejercicios.
    """
    nombre_mes = MESES_ES_COMPLETOS[inicio_mes.month]
    return f"{nombre_mes}_{inicio_mes.year}"


def _periodo_desde_nombre_hoja(sheet_name: str) -> Optional[tuple[int, int]]:
    """Reconoce hojas mensuales con formato Mes_AAAA."""
    for month_number, month_name in MESES_ES_COMPLETOS.items():
        match = re.fullmatch(
            rf"{re.escape(month_name)}_(\d{{4}})",
            sheet_name,
            flags=re.IGNORECASE,
        )
        if match:
            return int(match.group(1)), month_number

    return None


def _leer_indice_existente(workbook) -> dict[str, dict]:
    """
    Recupera metadatos del índice anterior para no perder la fecha de
    actualización de hojas antiguas que ya no se recalculan.
    """
    if SHEET_INDICE_MESES not in workbook.sheetnames:
        return {}

    worksheet = workbook[SHEET_INDICE_MESES]
    rows = list(worksheet.iter_rows(values_only=True))

    if not rows:
        return {}

    headers = [str(value or "").strip() for value in rows[0]]
    index_by_name = {name: index for index, name in enumerate(headers)}

    if "hoja" not in index_by_name:
        return {}

    result = {}

    for row in rows[1:]:
        sheet_name = row[index_by_name["hoja"]]
        if not sheet_name:
            continue

        item = {}
        for field, index in index_by_name.items():
            item[field] = row[index] if index < len(row) else ""

        result[str(sheet_name)] = item

    return result


def _construir_indice_meses(
    workbook,
    actualizados: list[dict],
    fecha_actualizacion: str,
) -> pd.DataFrame:
    """
    Genera un catálogo de todas las hojas mensuales disponibles.

    Incluye tanto los meses recalculados en la ejecución actual como las hojas
    antiguas que fueron conservadas desde ejecuciones previas.
    """
    indice_anterior = _leer_indice_existente(workbook)
    updated_by_sheet = {item["sheet"]: item for item in actualizados}
    current_period = actualizados[0]["periodo"] if actualizados else ""
    rows = []

    for sheet_name in workbook.sheetnames:
        parsed = _periodo_desde_nombre_hoja(sheet_name)
        if parsed is None:
            continue

        year, month = parsed
        period = f"{year:04d}-{month:02d}"
        month_start = datetime(
            year,
            month,
            1,
            OPERATIONAL_DAY_START_HOUR,
            0,
            0,
            tzinfo=TZ_CHILE,
        )
        month_end = add_months(month_start, 1)
        existing = indice_anterior.get(sheet_name, {})
        updated = updated_by_sheet.get(sheet_name)
        worksheet = workbook[sheet_name]

        if updated is not None:
            range_start = updated["from"]
            range_end = updated["to"]
            updated_at = fecha_actualizacion
        else:
            range_start = month_start
            range_end = month_end
            updated_at = existing.get("fecha_actualizacion", "")

        rows.append({
            "periodo": period,
            "anio": year,
            "mes_numero": month,
            "mes": MESES_ES_COMPLETOS[month],
            "hoja": sheet_name,
            "estado": "En curso" if period == current_period else "Cerrado",
            "desde": range_start.strftime("%Y-%m-%d %H:%M:%S"),
            "hasta": range_end.strftime("%Y-%m-%d %H:%M:%S"),
            "filas": max(worksheet.max_row - 1, 0),
            "fecha_actualizacion": updated_at,
        })

    rows.sort(key=lambda item: item["periodo"], reverse=True)

    columns = [
        "periodo",
        "anio",
        "mes_numero",
        "mes",
        "hoja",
        "estado",
        "desde",
        "hasta",
        "filas",
        "fecha_actualizacion",
    ]
    return pd.DataFrame(rows, columns=columns)



def _ordenar_hojas_libro(workbook) -> None:
    """Ordena hojas de control y meses desde el más reciente al más antiguo."""
    priority_names = [
        SHEET_INDICE_MESES,
        SHEET_RESUMEN_AGUA,
        SHEET_TENDENCIA_AGUA,
    ]
    priority = [
        workbook[name]
        for name in priority_names
        if name in workbook.sheetnames
    ]

    monthly = []
    others = []

    for worksheet in workbook.worksheets:
        if worksheet.title in priority_names:
            continue

        parsed = _periodo_desde_nombre_hoja(worksheet.title)
        if parsed is None:
            others.append(worksheet)
        else:
            year, month = parsed
            monthly.append((year, month, worksheet))

    monthly.sort(key=lambda item: (item[0], item[1]), reverse=True)
    workbook._sheets = priority + [item[2] for item in monthly] + others

def guardar_mapa_agua_excel(
    path: Path = DEFAULT_OUTPUT_XLSX,
    now: Optional[datetime] = None,
) -> dict:
    """
    Actualiza MapaAgua.xlsx sin borrar el historial mensual acumulado.

    Comportamiento, usando agosto de 2026 como ejemplo:
    - Agua_Resumen se actualiza.
    - Agua_Tendencia se actualiza con agosto en curso.
    - Agosto_2026 se crea o reemplaza y se alimenta durante el mes.
    - Julio_2026, Junio_2026 y Mayo_2026 se recalculan.
    - Las hojas de meses más antiguos se conservan sin ser eliminadas.

    Al cambiar a septiembre, se crea Septiembre_2026 y Agosto_2026 queda como
    mes cerrado. De esta forma el libro crece mes a mes.
    """
    if now is None:
        now = datetime.now(TZ_CHILE)

    resumen = calcular_resumen_agua(now=now).loc[:, COLUMNAS_REPORTE_AGUA]
    tendencia_actual = calcular_tendencia_agua(
        now=now,
        months_back=0,
    ).loc[:, COLUMNAS_TENDENCIA_AGUA]

    fecha_actualizacion = datetime.now(TZ_CHILE).strftime("%Y-%m-%d %H:%M:%S")
    meses_actualizados = []

    # Incluye el mes actual (0) y los meses cerrados configurados (1..N).
    for months_back in range(0, MESES_HISTORICOS + 1):
        inicio_mes, fin_mes = get_operational_month_window(
            now=now,
            months_back=months_back,
        )
        dataframe = calcular_tendencia_agua(
            now=now,
            months_back=months_back,
        ).loc[:, COLUMNAS_TENDENCIA_AGUA]

        meses_actualizados.append({
            "months_back": months_back,
            "sheet": _nombre_hoja_mes(inicio_mes),
            "periodo": inicio_mes.strftime("%Y-%m"),
            "estado": "En curso" if months_back == 0 else "Cerrado",
            "from": inicio_mes,
            "to": fin_mes,
            "dataframe": dataframe,
        })

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temp_path = path.with_name(f"{path.stem}.tmp{path.suffix}")

    if temp_path.exists():
        temp_path.unlink()

    existing_file = path.exists()

    try:
        if existing_file:
            # Trabajar sobre una copia permite conservar hojas antiguas y hacer
            # el reemplazo final de manera atómica.
            shutil.copy2(path, temp_path)
            writer_options = {
                "engine": "openpyxl",
                "mode": "a",
                "if_sheet_exists": "replace",
            }
        else:
            writer_options = {
                "engine": "openpyxl",
                "mode": "w",
            }

        with pd.ExcelWriter(temp_path, **writer_options) as writer:
            # Elimina las hojas relativas de la versión anterior para evitar
            # confusión. Sus períodos se vuelven a generar con nombre estable.
            for legacy_sheet in (
                "Mes_Anterior",
                "Mes_Anterior_2",
                "Mes_Anterior_3",
            ):
                if legacy_sheet in writer.book.sheetnames:
                    del writer.book[legacy_sheet]

            resumen.to_excel(
                writer,
                sheet_name=SHEET_RESUMEN_AGUA,
                index=False,
            )
            tendencia_actual.to_excel(
                writer,
                sheet_name=SHEET_TENDENCIA_AGUA,
                index=False,
            )

            _ajustar_hoja_excel(
                writer.book[SHEET_RESUMEN_AGUA],
                resumen,
            )
            _ajustar_hoja_excel(
                writer.book[SHEET_TENDENCIA_AGUA],
                tendencia_actual,
            )

            for item in meses_actualizados:
                dataframe = item["dataframe"]
                sheet_name = item["sheet"]

                dataframe.to_excel(
                    writer,
                    sheet_name=sheet_name,
                    index=False,
                )
                _ajustar_hoja_excel(
                    writer.book[sheet_name],
                    dataframe,
                )

            indice = _construir_indice_meses(
                writer.book,
                meses_actualizados,
                fecha_actualizacion,
            )
            indice.to_excel(
                writer,
                sheet_name=SHEET_INDICE_MESES,
                index=False,
            )
            _ajustar_hoja_excel(
                writer.book[SHEET_INDICE_MESES],
                indice,
            )

            _ordenar_hojas_libro(writer.book)

        temp_path.replace(path)

    finally:
        if temp_path.exists():
            temp_path.unlink()

    updated_sheet_names = {item["sheet"] for item in meses_actualizados}
    all_monthly_sheets = indice["hoja"].tolist() if not indice.empty else []
    preserved_sheets = [
        name for name in all_monthly_sheets
        if name not in updated_sheet_names
    ]

    return {
        "path": path,
        "resumen_sheet": SHEET_RESUMEN_AGUA,
        "tendencia_sheet": SHEET_TENDENCIA_AGUA,
        "indice_sheet": SHEET_INDICE_MESES,
        "resumen_rows": len(resumen),
        "tendencia_rows": len(tendencia_actual),
        "monthly_sheets": [
            {
                "sheet": item["sheet"],
                "periodo": item["periodo"],
                "estado": item["estado"],
                "rows": len(item["dataframe"]),
            }
            for item in meses_actualizados
        ],
        "preserved_sheets": preserved_sheets,
        "total_monthly_sheets": len(all_monthly_sheets),
    }
